from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.page import PageMargins

wb = Workbook()
ws = wb.active
ws.title = "請求書"

# ── 列幅 (A margin | B date | C content | D qty | E unit | F tax | G amount | H margin)
for col, w in {"A":2,"B":14,"C":34,"D":7,"E":7,"F":11,"G":15,"H":2}.items():
    ws.column_dimensions[col].width = w

# ── 行高
heights = {
    1:6,   2:18,  3:32,  4:20,  5:18,  6:18,  7:8,   8:6,
    9:16,  10:20, 11:14, 12:16, 13:6,  14:30, 15:6,  16:20,
    17:6,  18:20, 19:20, 20:20, 21:6,  22:16, 23:20, 24:20,
    25:20, 26:20, 27:20, 28:6,  29:16, 30:6,
}
for r, h in heights.items():
    ws.row_dimensions[r].height = h

# ── カラー定数
DARK    = "0A0818"
DARK2   = "1A1245"
PURPLE  = "764BA2"
ACCENT  = "667EEA"
ACCENT2 = "A5B4FC"
LIGHT   = "F5F2FC"
LIGHT2  = "FAF8FF"
WHITE   = "FFFFFF"
GRAY    = "999999"
GRAY2   = "CCCCCC"
TEXT    = "1A1A2E"
BORDER  = "E2D9F3"

def fill(c): return PatternFill("solid", fgColor=c)
def side(style="thin", color=BORDER): return Side(style=style, color=color)
def no_side(): return Side(style=None)

def border(t=False, b=False, l=False, r=False, color=BORDER, style="thin"):
    s = side(style, color)
    n = no_side()
    return Border(top=s if t else n, bottom=s if b else n,
                  left=s if l else n, right=s if r else n)

def cell(row, col, val="", bold=False, sz=10, color=TEXT,
         bg=None, ha="left", va="center", wrap=False,
         brd=None, fmt=None, italic=False):
    c = ws.cell(row=row, column=col, value=val)
    c.font = Font(name="游ゴシック", bold=bold, size=sz, color=color, italic=italic)
    c.alignment = Alignment(horizontal=ha, vertical=va, wrap_text=wrap)
    if bg:  c.fill = fill(bg)
    if brd: c.border = brd
    if fmt: c.number_format = fmt
    return c

def mg(r1,c1,r2,c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

# col index helpers
A,B,C,D,E,F,G,H = 1,2,3,4,5,6,7,8

# ════════════════════════════════════════
# ヘッダー背景 (rows 1-7)
# ════════════════════════════════════════
for r in range(1, 8):
    for c in range(1, 9):
        ws.cell(row=r, column=c).fill = fill(DARK)

# ── 左：タイトル・番号・日付 ──────────────
# "Invoice" サブラベル
mg(2,B, 2,D)
cell(2,B, "INVOICE", sz=8, color=ACCENT2, bg=DARK, ha="left", va="center",
     italic=False, bold=True)

# "請 求 書" 大見出し
mg(3,B, 3,D)
cell(3,B, "請 求 書", bold=True, sz=22, color=WHITE, bg=DARK, va="center")

# 請求書番号
mg(5,B, 5,B)
cell(5,B, "No.", sz=8, color=ACCENT2, bg=DARK, va="center", bold=True)
mg(5,C, 5,D)
cell(5,C, "INV-20260525-001", sz=10, color=WHITE, bg=DARK, va="center", bold=True)

# 発行日
mg(6,B, 6,B)
cell(6,B, "DATE", sz=8, color=ACCENT2, bg=DARK, va="center", bold=True)
mg(6,C, 6,D)
cell(6,C, "2026年5月25日", sz=10, color=WHITE, bg=DARK, va="center")

# ── 右：発行者カード ──────────────────────
mg(2,F, 2,G)
cell(2,F, "ISSUED BY", sz=7, color=ACCENT2, bg=DARK, ha="right", va="center", bold=True)

mg(3,F, 3,G)
cell(3,F, "可知 優也", bold=True, sz=16, color=WHITE, bg=DARK, ha="right", va="center")

mg(4,F, 4,G)
cell(4,F, "愛知県名古屋市昭和区壇溪通1丁目7番地", sz=8, color=GRAY2, bg=DARK, ha="right", va="center")

mg(5,F, 5,G)
cell(5,F, "アンシェリーナ106", sz=8, color=GRAY2, bg=DARK, ha="right", va="center")

mg(6,F, 6,G)
cell(6,F, "TEL: 080-6925-9411", sz=8, color=GRAY2, bg=DARK, ha="right", va="center")

mg(7,F, 7,G)
cell(7,F, "免税事業者 · 登録番号なし", sz=7, color=ACCENT2, bg=DARK, ha="right", va="center")

# ── アクセントライン (row 7 bottom) ──
ws.row_dimensions[7].height = 4
for c in range(1, 9):
    ws.cell(row=7, column=c).fill = fill(PURPLE)

# ════════════════════════════════════════
# 宛先 (rows 9-11)
# ════════════════════════════════════════
mg(9,B, 9,G)
cell(9,B, "BILL TO  /  請求先", sz=7, color=PURPLE, ha="left", va="center",
     bold=True, bg=WHITE)

mg(10,B, 10,E)
c = cell(10,B, "合同会社 Frigus　御中", bold=True, sz=17, color=TEXT, bg=WHITE, va="bottom")
c.border = Border(bottom=Side(style="medium", color=DARK2))

mg(11,B, 11,G)
cell(11,B, "", bg=WHITE)

# ════════════════════════════════════════
# 請求金額 (row 12)
# ════════════════════════════════════════
ws.row_dimensions[12].height = 36
for c in range(1, 9):
    ws.cell(row=12, column=c).fill = fill(DARK2)

mg(12,B, 12,D)
cell(12,B, "ご請求金額（税込）", sz=9, color="8899CC", bg=DARK2, va="center")

mg(12,E, 12,G)
cell(12,E, 202306, bold=True, sz=20, color=WHITE, bg=DARK2, ha="right", va="center",
     fmt='¥#,##0')

# ════════════════════════════════════════
# 明細ヘッダー (row 14)
# ════════════════════════════════════════
hdr_brd = border(t=True, b=True, l=True, r=True, color=BORDER)

cell(14,B, "日付", bold=True, sz=8, color=PURPLE, bg=LIGHT, ha="center", brd=hdr_brd)

mg(14,C, 14,C)
cell(14,C, "内容", bold=True, sz=8, color=PURPLE, bg=LIGHT, brd=hdr_brd)

cell(14,D, "数量", bold=True, sz=8, color=PURPLE, bg=LIGHT, ha="center", brd=hdr_brd)
cell(14,E, "単位", bold=True, sz=8, color=PURPLE, bg=LIGHT, ha="center", brd=hdr_brd)
cell(14,F, "税区分", bold=True, sz=8, color=PURPLE, bg=LIGHT, ha="center", brd=hdr_brd)
cell(14,G, "金額", bold=True, sz=8, color=PURPLE, bg=LIGHT, ha="right", brd=hdr_brd)

# ════════════════════════════════════════
# 明細データ (row 16)
# ════════════════════════════════════════
dat_brd = border(t=True, b=True, l=True, r=True)
ws.row_dimensions[16].height = 26

cell(16,B, "2026年5月25日", sz=9, color=GRAY, ha="center", brd=dat_brd)
cell(16,C, "Goo Property Singapore Pte. Ltd.社からの保全費（2026年2月預かり金より）",
     sz=9, color=TEXT, wrap=True, brd=dat_brd)
cell(16,D, 1, sz=9, color=TEXT, ha="center", brd=dat_brd)
cell(16,E, "式", sz=9, color=TEXT, ha="center", brd=dat_brd)
cell(16,F, "対象外", sz=9, color=PURPLE, ha="center",
     bg="EDE9FB", brd=dat_brd)
cell(16,G, 202306, sz=9, color=TEXT, ha="right", brd=dat_brd, fmt='¥#,##0')

# ════════════════════════════════════════
# 合計 (rows 18-20)
# ════════════════════════════════════════
for row, label, val, total in [
    (18, "小計",   "¥202,306", False),
    (19, "消費税", "¥0（対象外）", False),
    (20, "合　計", 202306,    True),
]:
    mg(row, E, row, F)
    cell(row, E, label,
         bold=total, sz=9 if not total else 11,
         color=WHITE if total else GRAY,
         bg=DARK2 if total else WHITE, ha="right",
         brd=border(t=True, b=True, l=True, r=True,
                    color=BORDER if not total else DARK2))
    c = ws.cell(row=row, column=G,
                value=val if not total else 202306)
    c.font = Font(name="游ゴシック", bold=total,
                  size=11 if total else 9,
                  color=WHITE if total else GRAY)
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.fill = fill(DARK2) if total else fill(WHITE)
    if total: c.number_format = '¥#,##0'
    c.border = border(t=True, b=True, l=True, r=True,
                      color=BORDER if not total else DARK2)

# ════════════════════════════════════════
# 振込先 (rows 22-27)
# ════════════════════════════════════════
mg(22,B, 22,G)
cell(22,B, "BANK TRANSFER  /  お振込先", sz=7, color=PURPLE, bold=True)

bank = [
    (23, "銀行名",   "楽天銀行"),
    (24, "支店名",   "ボレロ支店"),
    (25, "口座種別", "普通預金"),
    (26, "口座番号", "4541728"),
    (27, "口座名義", "カチユウヤ"),
]
bk_brd = border(t=True, b=True, l=True, r=True)
for row, key, val in bank:
    cell(row, B, key, bold=True, sz=8, color=GRAY, bg=LIGHT, brd=bk_brd)
    mg(row, C, row, G)
    cell(row, C, val, sz=10, color=DARK2, bg=LIGHT2, brd=bk_brd)

# ════════════════════════════════════════
# フッター注記 (row 29)
# ════════════════════════════════════════
mg(29,B, 29,G)
cell(29,B, "ご不明な点はお気軽にご連絡ください。  TEL: 080-6925-9411",
     sz=8, color=GRAY2, italic=True)

# ════════════════════════════════════════
# 印刷設定
# ════════════════════════════════════════
ws.page_setup.orientation = "portrait"
ws.page_setup.paperSize   = 9
ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
ws.print_area = "A1:H30"

wb.save("/home/user/test/請求書_INV-20260525-001.xlsx")
print("Done")
